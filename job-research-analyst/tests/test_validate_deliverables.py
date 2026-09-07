#!/usr/bin/env python3
"""Structure tests for validate_deliverables.py (P4 / P4b / N2–N5)."""
from __future__ import annotations

import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "scripts" / "validate_deliverables.py"
TEMPLATES = ROOT / "assets" / "templates"
REPORT = ROOT / "references" / "report-templates.md"
sys.path.insert(0, str(ROOT / "scripts"))
import validate_deliverables as vd  # noqa: E402


def run_validator(args: list[str]) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, str(SCRIPT), *args],
        capture_output=True,
        text=True,
        encoding="utf-8",
    )


def named_copies(tmpdir: Path) -> tuple[Path, Path, Path]:
    docx = tmpdir / "星河数科_深度介绍_20260906.docx"
    salary = tmpdir / "政务项目经理_薪资福利对比_20260906.xlsx"
    review = tmpdir / "星河数科_员工评价与风险提示_20260906.xlsx"
    shutil.copy2(TEMPLATES / "company-profile.docx", docx)
    shutil.copy2(TEMPLATES / "salary-benefits.xlsx", salary)
    shutil.copy2(TEMPLATES / "reviews-risks.xlsx", review)
    return docx, salary, review


class TestFrozenSync(unittest.TestCase):
    def test_p4b_constants_in_report_templates(self) -> None:
        text = REPORT.read_text(encoding="utf-8")
        missing = [s for s in vd.FROZEN_STRINGS if s not in text]
        self.assertEqual(missing, [], f"frozen strings missing from report-templates.md: {missing}")


class TestEmptySkeleton(unittest.TestCase):
    def test_p4_empty_skeleton_files(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            tmp = Path(td)
            docx, salary, review = named_copies(tmp)
            r = run_validator(["--files", str(docx), str(salary), str(review)])
            self.assertEqual(r.returncode, 0, r.stdout + r.stderr)

    def test_p5_dir_mode(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            tmp = Path(td)
            named_copies(tmp)
            r = run_validator(["--dir", str(tmp)])
            self.assertEqual(r.returncode, 0, r.stdout + r.stderr)


class TestNegatives(unittest.TestCase):
    def test_n2_missing_risk_sheet(self) -> None:
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as td:
            tmp = Path(td)
            _, _, review = named_copies(tmp)
            wb = load_workbook(review)
            wb.remove(wb["风险提示清单"])
            wb.save(review)
            r = run_validator(["--files", str(review)])
            self.assertNotEqual(r.returncode, 0)
            self.assertIn("风险提示清单", r.stdout)

    def test_n3_bad_filename(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            tmp = Path(td)
            bad = tmp / "intro.docx"
            shutil.copy2(TEMPLATES / "company-profile.docx", bad)
            r = run_validator(["--files", str(bad)])
            self.assertNotEqual(r.returncode, 0)

    def test_n4_missing_heading(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            tmp = Path(td)
            docx, _, _ = named_copies(tmp)
            with zipfile.ZipFile(docx, "r") as zf:
                xml = zf.read("word/document.xml").decode("utf-8")
                others = {n: zf.read(n) for n in zf.namelist() if n != "word/document.xml"}
            xml = xml.replace("岗位 JD 解读", "岗位说明")
            xml = xml.replace("信息来源清单", "参考资料")
            out = tmp / "patched.docx"
            with zipfile.ZipFile(out, "w") as zf:
                zf.writestr("word/document.xml", xml)
                for n, data in others.items():
                    zf.writestr(n, data)
            renamed = tmp / "星河数科_深度介绍_20260906.docx"
            out.replace(renamed)
            r = run_validator(["--files", str(renamed)])
            self.assertNotEqual(r.returncode, 0)
            self.assertTrue("岗位 JD 解读" in r.stdout or "信息来源清单" in r.stdout)

    def test_n5_missing_review_subject_column(self) -> None:
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as td:
            tmp = Path(td)
            _, _, review = named_copies(tmp)
            wb = load_workbook(review)
            ws = wb["员工评价汇总"]
            ws.cell(1, 6).value = "备注"
            wb.save(review)
            r = run_validator(["--files", str(review)])
            self.assertNotEqual(r.returncode, 0)
            self.assertIn("评价主体", r.stdout)


class TestDescriptionLength(unittest.TestCase):
    def test_skill_description_le_1024(self) -> None:
        import re

        text = (ROOT / "SKILL.md").read_text(encoding="utf-8")
        m = re.search(r"^description:\s*>\s*\n((?:  .*\n)+)", text, re.M)
        self.assertIsNotNone(m)
        desc = " ".join(line.strip() for line in m.group(1).splitlines() if line.strip())
        self.assertLessEqual(len(desc), 1024, len(desc))
        self.assertIn("查一下这个公司", desc)


if __name__ == "__main__":
    unittest.main()
